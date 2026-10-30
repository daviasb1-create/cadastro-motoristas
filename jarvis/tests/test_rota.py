"""Roteirização: motor (regras da ferramenta), leitura da pasta, ações e comandos de voz.

Usa a configuração de exemplo (motoristas fictícios); a configuração real da
filial fica só no PC do usuário.
"""
import json
import os
import time
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

from jarvis.actions.base import REGISTRY
from jarvis.integrations.rota import service as rota_service
from jarvis.integrations.rota.engine import FilialConfig, Report, RouteEngine, parse_date, parse_num
from jarvis.integrations.rota.reader import FOLDERS
from jarvis.nlu.intents import IntentParser

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = json.loads((ROOT / "config" / "rota_filial.example.json").read_text(encoding="utf-8"))
HEAD = ["Notas Fiscais", "Região de Entrega", "CEP Destinatário", "Peso Real", "Peso Cubado", "Prazo Entrega",
        "Remetente", "Cidade Destinatário", "Bairro Destinatário", "Endereço Destinatário", "Valor do Frete", "Ocorrência"]
D = date(2026, 11, 2)


def row(nf, reg="RIO-OES1", cep="22000000", peso=5, cub=5, prazo="05/11/2026", cliente="LOJA X", cidade="Rio de Janeiro",
        bairro="Centro", frete=100, occ=""):
    return [nf, reg, cep, peso, cub, prazo, cliente, cidade, bairro, "Rua A, 1", frete, occ]


def report(rows):
    return Report(rows=[dict(zip(HEAD, r)) for r in rows], headers=list(HEAD))


def engine(**changes):
    cfg = dict(EXAMPLE)
    cfg.update(changes)
    return RouteEngine(FilialConfig.from_dict(cfg))


def by_driver(result):
    return {m.driver.motorista: [n.nf for n in m.items] for m in result.manifests}


def test_parse_helpers():
    assert parse_num("1.234,5") == 1234.5
    assert parse_num("12,5 kg") == 12.5
    assert parse_num("") != parse_num("")  # NaN
    assert parse_date("05/11/26") == date(2026, 11, 5)
    assert parse_date("2026-11-05 00:00:00") == date(2026, 11, 5)
    assert parse_date(46331) == date(2026, 11, 5)


def test_region_cep_and_bairro_rules():
    e = engine()
    rep = report([
        row("1", reg="RIO-OES1"),
        row("2", reg="", cep="24400000", cidade="São Gonçalo"),  # CEP 244 -> NITKY mesmo sem região
        row("3", reg="", bairro="Barra da Tijuca"),  # Barra no Rio -> OES 5
        row("4", reg="", bairro="Barra da Tijuca", cidade="Nova Iguaçu", cep="26000000"),  # homônimo fora da capital
        row("5", reg="RIO-NVIGU"),  # região sem motorista fixo
    ])
    r = e.assign({"no_rj": rep}, D)
    got = by_driver(r)
    assert got["CARLOS EXEMPLO"] == ["1"]
    assert got["PAULO EXEMPLO"] == ["2"]
    assert got["RAFAEL EXEMPLO"] == ["3"]
    reasons = {n.nf: n.reason for n in r.overflow}
    assert reasons["4"] == "Sem região"
    assert "sem motorista fixo" in reasons["5"]
    dois = next(m for m in r.manifests if m.driver.motorista == "PAULO EXEMPLO")
    assert dois.custo == 380 + 50 and dois.addon_label == "São Gonçalo"


def test_priority_delivered_and_duplicates():
    e = engine()
    chegando = report([row("10", cliente="CLIENTE PRIORITARIO LTDA", prazo="30/11/2026"),
                       row("11", cliente="CLIENTE PRIORITARIO EXCECAO", prazo="30/11/2026"),
                       row("12", occ="Entrega realizada"),
                       row("13", prazo="30/11/2026")])
    agendado = report([row("13", prazo="30/11/2026")])
    r = e.assign({"chegando": chegando, "agendado": agendado}, D)
    notes = {n.nf: n for m in r.manifests for n in m.items}
    assert notes["10"].prioridade and not notes["11"].prioridade
    assert "12" not in notes and r.iss.entregue == 1
    assert r.iss.dup == 1 and notes["13"].agendado  # duplicada mantém o agendamento
    assert r.must_ship(notes["13"]) and not r.must_ship(notes["11"])


def test_capacity_heavy_and_mandatory_first():
    e = engine()
    rows = [row(str(100 + i), peso=10, prazo="30/11/2026") for i in range(30)]
    rows.append(row("200", peso=40, prazo="30/11/2026"))  # pesado não obrigatório: vai por último na Fiorino
    rows.append(row("201", peso=10, prazo="01/11/2026"))  # carga do dia: entra primeiro
    r = e.assign({"no_rj": report(rows)}, D)
    um = next(m for m in r.manifests if m.driver.motorista == "CARLOS EXEMPLO")
    assert len(um.items) == 30 and "201" in [n.nf for n in um.items]
    fora = {n.nf: n.reason for n in r.overflow}
    assert fora["200"].startswith("Peso elevado")
    assert sum(1 for v in fora.values() if v.startswith("Região cheia")) == 1
    assert um.recomenda in ("ok", "tetoprazo", "segurar")


def test_cost_band_and_hold_recommendation():
    e = engine()
    r = e.assign({"no_rj": report([row("1", frete=500, prazo="30/11/2026")])}, D)
    um = next(m for m in r.manifests if m.driver.motorista == "CARLOS EXEMPLO")
    assert um.pct_custo == pytest.approx(350 / 500) and um.band == "vermelho" and um.recomenda == "segurar"
    r = e.assign({"no_rj": report([row("1", frete=500, prazo="01/11/2026")])}, D)
    um = next(m for m in r.manifests if m.driver.motorista == "CARLOS EXEMPLO")
    assert um.recomenda == "tetoprazo"
    r = e.assign({"no_rj": report([row("1", frete=2000)])}, D)
    assert next(m for m in r.manifests if m.driver.motorista == "CARLOS EXEMPLO").band == "verde"


def test_absent_removed_and_reassign():
    e = engine()
    rep = report([row("1"), row("2", reg="RIO-NVIGU"), row("3", reg="", bairro="Barra da Tijuca")])
    r = e.assign({"no_rj": rep}, D, absent={"CARLOS EXEMPLO"}, removed={"3"})
    um = next(m for m in r.manifests if m.driver.motorista == "CARLOS EXEMPLO")
    assert um.absent and [n.nf for n in um.items] == ["1"]
    assert any(n.nf == "3" and n.removed for n in r.overflow)
    r = e.assign({"no_rj": rep}, D, reassign={"CARLOS EXEMPLO": "RIO-NVIGU"})
    got = by_driver(r)
    assert got["CARLOS EXEMPLO"] == ["2"]
    assert any(n.nf == "1" and n.realoc_fora for n in r.overflow)


# ----------------------------------------------------------- pasta e ações
def write_xlsx(path: Path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


@pytest.fixture
def rota(services, tmp_path, monkeypatch):
    cfg = tmp_path / "rota_filial.json"
    cfg.write_text(json.dumps(EXAMPLE), encoding="utf-8")
    monkeypatch.setattr(rota_service, "config_paths", lambda: [cfg])
    services.settings.set("rota.folder", str(tmp_path / "rotas"))
    services.extras["rota"] = rota_service.RotaService(services.settings, services.bus)
    return tmp_path / "rotas"


def test_build_by_voice_from_folder(services, rota):
    tomorrow = date.today() + timedelta(days=1)
    prazo = tomorrow.strftime("%d/%m/%Y")
    write_xlsx(rota / FOLDERS["no_rj"] / "cargas.xlsx", [row(1001, prazo=prazo), row(1002, reg="RIO-OES5", frete=3000),
                                                         row(1003, reg="RIO-NVIGU", prazo=prazo)])
    old = rota / FOLDERS["chegando"] / "ontem.xlsx"
    write_xlsx(old, [row(9999)])
    os.utime(old, (time.time() - 3 * 86400,) * 2)
    parser = IntentParser()
    intent = parser.parse("Jarvis, monta a rota de amanhã")
    assert intent.name == "rota_montar"
    result = REGISTRY.run(services, intent.name, intent.params)
    assert result.ok, result.say
    assert "2 notas em 2 rotas" in result.say and "1 carga obrigatória sem cobertura" in result.say
    assert "ontem.xlsx" in result.display  # planilha antiga ignorada e avisada
    svc = services.extras["rota"]
    assert svc.excel_path.exists() and svc.print_path.exists()
    wb = openpyxl.load_workbook(svc.excel_path)
    assert wb.sheetnames == ["Resumo", "Romaneio", "Fora realocar"]
    assert wb["Romaneio"]["F2"].value in ("1001", "1002")

    intent = parser.parse("rota do carlos")
    assert intent.name == "rota_motorista"
    say = REGISTRY.run(services, intent.name, intent.params).say
    assert say.startswith("Carlos, OES 1") and "1 nota," in say

    assert parser.parse("o que ficou fora da rota").name == "rota_fora"
    assert "1 sem motorista fixo" in REGISTRY.run(services, "rota_fora", {}).say

    intent = parser.parse("tira a nota 1002")
    assert intent.name == "rota_remover_nota"
    res = REGISTRY.run(services, intent.name, intent.params)
    assert res.ok and "retirada" in res.say and "1 nota em 1 rota" in res.say

    res = REGISTRY.run(services, "rota_realocar", {"driver": "rafael", "region": "nvigu"})
    assert res.ok and "RIO-NVIGU" in res.say and "Toda a carga obrigatória está coberta" in res.say

    res = REGISTRY.run(services, "rota_ausente", {"driver": "carlos"})
    assert res.ok and "ausente" in res.say

    assert REGISTRY.run(services, "rota_abrir", {"what": "romaneio"}).ok
    assert services.opened[-1].endswith(".html")


def test_missing_sheets_and_config(services, rota, monkeypatch):
    res = REGISTRY.run(services, "rota_montar", {"day": "hoje"})
    assert not res.ok and "Não encontrei planilhas" in res.say
    monkeypatch.setattr(rota_service, "config_paths", lambda: [rota / "nao_existe.json"])
    write_xlsx(rota / FOLDERS["no_rj"] / "cargas.xlsx", [row(1)])
    res = REGISTRY.run(services, "rota_montar", {"day": "hoje"})
    assert not res.ok and "rota_filial.json" in res.say


@pytest.mark.parametrize("phrase,name,params", [
    ("faz a roteirização de hoje", "rota_montar", None),
    ("como ficou a rota", "rota_resumo", {}),
    ("o elias faltou", "rota_ausente", {"driver": "elias"}),
    ("o elias voltou", "rota_ausente", {"driver": "elias", "present": True}),
    ("realoca o elias para oes seis", "rota_realocar", {"driver": "elias", "region": "oes 6"}),
    ("o marcos vai fazer a dcx", "rota_realocar", {"driver": "marcos", "region": "dcx"}),
    ("imprime os romaneios", "rota_abrir", {"what": "romaneio"}),
    ("abre a planilha da rota", "rota_abrir", {"what": "excel"}),
    ("restaura a nota 12345", "rota_remover_nota", {"nf": "12345", "restore": True}),
])
def test_rota_phrases(phrase, name, params):
    intent = IntentParser().parse(phrase)
    assert intent is not None and intent.name == name, (phrase, intent)
    if params is not None:
        assert intent.params == params


def test_outlook_disabled_by_default(services):
    from jarvis.integrations.msgraph.auth import GraphAuth

    auth = GraphAuth(services.settings)
    assert "Mail.ReadWrite" not in auth.scopes and "Files.ReadWrite.All" in auth.scopes
    services.settings.set("outlook.enabled", True)
    assert "Mail.ReadWrite" in auth.scopes
