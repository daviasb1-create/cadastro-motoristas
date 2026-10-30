"""Saídas da rota: Excel (Resumo, Romaneio, Fora realocar) e romaneios para impressão.

Mesmas colunas da ferramenta original, para o arquivo continuar compatível com
quem já usa o Excel exportado pelo navegador.
"""
from __future__ import annotations

import html
from pathlib import Path

from .engine import Manifest, RouteResult, _fmt


def fmt_date(d) -> str:
    return d.strftime("%d/%m/%Y") if d else "–"


def fmt_cep(c: str) -> str:
    return (c[:5] + "-" + c[5:] if len(c) > 5 else c) if c else "–"


def situacao(m: Manifest, r: RouteResult) -> str:
    if not m.band:
        return ""
    lim, ver = round(r.limite * 100), round(r.vermelho * 100)
    band = {"verde": f"Dentro (≤{lim}%)", "bege": f"Bege ({lim}-{ver}%)"}.get(m.band, f"Vermelho (>{ver}%)")
    extra = " · pode segurar" if m.recomenda == "segurar" else " · vai obrigatório" if m.recomenda == "tetoprazo" else ""
    return band + extra


def write_excel(result: RouteResult, path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Motorista", "Turno", "Viagem", "Veículo", "Região", "Disponível", "NFs", "Máx NF", "Peso (kg)", "Peso máx",
               "Peso cubado (kg)", "Cubado máx", "Gargalo", "Valor frete (R$)", "Custo agregado", "Adicional",
               "% Custo/Frete", "Resultado", "Situação"])
    for m in result.manifests:
        d = m.driver
        ws.append([d.motorista, "", "", d.veiculo, d.regiao, "NÃO" if m.absent else "sim", len(m.items), d.maxnf,
                   round(m.peso, 1), d.peso, round(m.vol, 1), d.vol,
                   (m.gargalo if m.mx >= 0.85 else "com folga") if m.items else "—", round(m.valor_frete, 2), m.custo,
                   m.addon_label, f"{round(m.pct_custo * 100)}%" if m.pct_custo is not None else "",
                   round(m.resultado, 2), situacao(m, result) if m.items else ""])
    rom = wb.create_sheet("Romaneio")
    rom.append(["Motorista", "Turno", "Veículo", "Região", "Disponível", "NF", "CT-e", "Cliente", "CEP", "Município",
                "Bairro", "Endereço", "Peso (kg)", "Peso cubado (kg)", "Valor frete (R$)", "Prazo", "Carga do dia", "Origem"])
    for m in sorted(result.manifests, key=lambda m: m.driver.seq):
        for r in m.items:
            rom.append([m.driver.motorista, "", m.driver.veiculo, m.driver.regiao, "NÃO" if m.absent else "sim", r.nf,
                        r.cte, r.cliente, r.cep, r.municipio, r.bairro, r.endereco, round(r.peso, 1), round(r.volume, 1),
                        round(r.valor_frete or 0, 2), fmt_date(r.prazo),
                        "SIM" if r.prazo and r.prazo <= result.data_rota else "", r.origem])
    ov = wb.create_sheet("Fora realocar")
    ov.append(["NF", "Cliente", "CEP", "Município", "Bairro", "Região", "Peso (kg)", "Peso cubado (kg)", "Prazo",
               "Carga do dia", "Motivo"])
    for r in sorted(result.overflow, key=lambda n: not n.urgent):  # obrigatórias primeiro, como na ferramenta
        ov.append([r.nf, r.cliente, r.cep, r.municipio, r.bairro, r.reg_label, round(r.peso, 1), round(r.volume, 1),
                   fmt_date(r.prazo), "SIM" if r.urgent else "", r.reason])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column[:200])
            sheet.column_dimensions[column[0].column_letter].width = min(45, max(8, width + 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_romaneios(result: RouteResult, path: Path, filial: str = "ÍCARO EXPRESS") -> Path:
    """HTML com um romaneio por página; abre no navegador já pedindo impressão."""
    e = html.escape
    hoje = fmt_date(result.data_rota)
    pages = []
    for m in sorted(result.manifests, key=lambda m: (m.driver.carreg, m.driver.seq)):
        if not m.items:
            continue
        d = m.driver
        rows = "".join(
            f"<tr><td>{i}</td><td>{e(r.nf)}</td><td>{e(r.cte)}</td><td>{e(r.cliente)}</td><td>{fmt_cep(r.cep)}</td>"
            f"<td>{e(' / '.join(x for x in (r.municipio, r.bairro) if x))}</td><td>{e(r.endereco)}</td>"
            f"<td>{_fmt(r.peso, 1)}</td><td>{fmt_date(r.prazo)}</td><td>&nbsp;&nbsp;</td></tr>"
            for i, r in enumerate(m.items, 1))
        pages.append(
            f"<div class='romaneio'><h3>{e(filial)} · Romaneio{' (REALOCAR: motorista ausente)' if m.absent else ''}</h3>"
            f"<div class='rmeta'>Motorista: <b>{e(d.motorista)}</b> | Veículo: {e(d.veiculo)} | Região: {e(d.regiao)} | "
            f"Carregamento {d.carreg} | {len(m.items)} NFs | {_fmt(m.peso)} kg | {_fmt(m.vol)} kg cub | Data: {hoje} | "
            f"Placa: __________</div><table><thead><tr><th>#</th><th>NF</th><th>CT-e</th><th>Cliente</th><th>CEP</th>"
            f"<th>Município / Bairro</th><th>Endereço</th><th>Peso</th><th>Prazo</th><th>OK</th></tr></thead>"
            f"<tbody>{rows}</tbody></table><div class='sig'><div>Assinatura do motorista</div>"
            f"<div>Conferente / Expedição</div></div></div>")
    doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Romaneios {hoje}</title>
<style>
body{{font-family:Segoe UI,Arial,sans-serif;font-size:11px;color:#111;margin:16px}}
.romaneio{{page-break-after:always;margin-bottom:24px}}
h3{{margin:0 0 4px;font-size:15px}} .rmeta{{margin-bottom:8px}}
table{{width:100%;border-collapse:collapse}} th,td{{border:1px solid #999;padding:3px 4px;text-align:left}}
th{{background:#eee}} .sig{{display:flex;gap:40px;margin-top:36px}} .sig div{{flex:1;border-top:1px solid #000;padding-top:4px;text-align:center}}
</style></head><body onload="window.print()">{''.join(pages) or '<p>Nenhuma rota com carga.</p>'}</body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(doc, encoding="utf-8")
    return path
